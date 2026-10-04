import json
import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

JSON_FILE = "risk_matrix.json"
EXCEL_FILE = "financial_risk_heatmap.xlsx"

if not os.path.exists(JSON_FILE):
    print(f"❌ Ошибка: Файл {JSON_FILE} еще не создан моделью!")
    exit()

with open(JSON_FILE, "r", encoding="utf-8") as f:
    data = json.load(f)

# ИСПРАВЛЕНИЕ ОПЕЧАТКИ ИИ: Корректируем имя второй компании в массиве
if len(data) >= 2:
    data[1]["Company"] = "Seligdar"

df = pd.DataFrame(data)

df.rename(columns={
    "Company": "Компания",
    "Interest_Rate_Risk": "Процентный риск (Ставка ЦБ)",
    "Currency_Risk": "Валютный риск (Курс рубля)",
    "Labor_Shortage_Risk": "Кадровый дефицит (ФОТ)",
    "Supply_Chain_Risk": "Логистика и Оборудование (CAPEX)"
}, inplace=True)

df.to_excel(EXCEL_FILE, index=False)
print(f"📊 Базовый файл {EXCEL_FILE} успешно сформирован.")

wb = load_workbook(EXCEL_FILE)
ws = wb.active

fills = {
    1: PatternFill(start_color="2ECC71", end_color="2ECC71", fill_type="solid"), # Яркий зеленый (Низкий риск)
    2: PatternFill(start_color="A9DFBF", end_color="A9DFBF", fill_type="solid"), # Салатовый
    3: PatternFill(start_color="F1C40F", end_color="F1C40F", fill_type="solid"), # Насыщенный желтый (Средний)
    4: PatternFill(start_color="E67E22", end_color="E67E22", fill_type="solid"), # Яркий оранжевый (Высокий)
    5: PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid")  # Насыщенный красный (Критический)
}

fonts = {
    1: Font(color="000000", bold=True),
    2: Font(color="000000", bold=True),
    3: Font(color="000000", bold=True),
    4: Font(color="FFFFFF", bold=True),
    5: Font(color="FFFFFF", bold=True)
}

for row in range(2, ws.max_row + 1):
    for col in range(2, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        try:
            val = int(cell.value)
            if val in fills:
                cell.fill = fills[val]
                cell.font = fonts[val]
        except (ValueError, TypeError):
            pass

# Исправленный блок автоподбора ширины колонок
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    # Берем букву колонки от её первой ячейки
    col_letter = col[0].column_letter
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

wb.save(EXCEL_FILE)
print(f"🎨 Тепловая карта успешно построена и сохранена в {EXCEL_FILE}!")